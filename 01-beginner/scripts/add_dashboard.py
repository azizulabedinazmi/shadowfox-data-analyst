"""
Add the 'Monthly Trend' and 'Dashboard' sheets to superstore_beginner.xlsx.

Run after build_workbook.py:  python scripts/add_dashboard.py
Set PREVIEW=1 to hide the data sheets and write superstore_preview.xlsx (used
only to render a picture of the dashboard).

Monthly Trend = formula-based summary table (SUMIFS on the Data sheet), the same
numbers a PivotTable would give, plus a seasonality grid.
Dashboard     = KPI cards linked to 'KPI Definitions' + charts.
"""
import os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "superstore_beginner.xlsx"
PREVIEW = os.environ.get("PREVIEW") == "1"
OUT = ROOT / ("superstore_preview.xlsx" if PREVIEW else "superstore_beginner.xlsx")

wb = load_workbook(SRC)
for name in ("Dashboard", "Monthly Trend"):
    if name in wb.sheetnames:
        del wb[name]

kpi = wb["KPI Definitions"]
data = wb["Data"]
last = data.max_row  # last data row
# Guard against the KPI sheet layout moving: the chart ranges below depend on it.
assert kpi["B29"].value == "Furniture" and kpi["B31"].value == "Technology"
assert kpi["B36"].value == "Central" and kpi["B39"].value == "West"
assert data["V1"].value == "Order Year" and data["X1"].value == "Year-Month"
pa = wb["Product Analysis"]
da = wb["Discount Analysis"]
assert pa["L5"].value == "Sub-Category" and pa["B28"].value == "Product (short)"
assert pa["B42"].value == "Product (short)" and da["G12"].value == "Profit Margin"

FONT = "Arial"
f = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)  # noqa: E731
NAVY, BLUE, LIGHT, GREY = "1F3864", "2E75B6", "DEEAF6", "F2F2F2"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
D = lambda col: f"Data!${col}$2:${col}${last}"  # noqa: E731


def head(ws, row, col, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_axes(ch):
    ch.x_axis.tickLblPos = "low"  # keep labels clear of negative values
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"


# =========================================================== Monthly Trend sheet
mt = wb.create_sheet("Monthly Trend", 0)
mt["A1"] = "Monthly sales trend, 2015-2018"
mt["A1"].font = f(size=14, bold=True)
mt["A2"] = ("Formula-based summary of the Data sheet (SUMIFS on the Year-Month column). "
            "Gross figures: returned orders are included.")
mt["A2"].font = f(size=9, italic=True, color="595959")

head(mt, 4, 1, ["Year-Month", "Year", "Month", "Sales", "Profit", "Profit Margin",
                "Orders", "Sales MoM %", "3-Month Avg Sales"])
years = [2015, 2016, 2017, 2018]
r = 5
for y in years:
    for m in range(1, 13):
        mt.cell(row=r, column=1, value=f'=B{r}&"-"&TEXT(C{r},"00")')
        mt.cell(row=r, column=2, value=y)
        mt.cell(row=r, column=3, value=m)
        mt.cell(row=r, column=4, value=f"=SUMIFS({D('R')},{D('X')},$A{r})")
        mt.cell(row=r, column=5, value=f"=SUMIFS({D('U')},{D('X')},$A{r})")
        mt.cell(row=r, column=6, value=f"=IF(D{r}=0,0,E{r}/D{r})")
        mt.cell(row=r, column=7, value=f"=SUMIFS({D('AC')},{D('X')},$A{r})")
        if r > 5:
            mt.cell(row=r, column=8, value=f"=IF(D{r-1}=0,0,D{r}/D{r-1}-1)")
        if r > 6:
            mt.cell(row=r, column=9, value=f"=AVERAGE(D{r-2}:D{r})")
        for c, nf in zip(range(1, 10), ["@", "0", "0", "$#,##0", "$#,##0", "0.0%",
                                        "#,##0", "0.0%", "$#,##0"]):
            cell = mt.cell(row=r, column=c)
            cell.font, cell.number_format, cell.border = f(), nf, box
            if (r % 2) == 0:
                cell.fill = PatternFill("solid", fgColor=GREY)
        r += 1
first, lastm = 5, r - 1  # 5..52
tot = r
mt.cell(row=tot, column=1, value="Total")
mt.cell(row=tot, column=4, value=f"=SUM(D{first}:D{lastm})")
mt.cell(row=tot, column=5, value=f"=SUM(E{first}:E{lastm})")
mt.cell(row=tot, column=6, value=f"=IF(D{tot}=0,0,E{tot}/D{tot})")
mt.cell(row=tot, column=7, value=f"=SUM(G{first}:G{lastm})")
for c, nf in [(1, "@"), (4, "$#,##0"), (5, "$#,##0"), (6, "0.0%"), (7, "#,##0")]:
    cell = mt.cell(row=tot, column=c)
    cell.font, cell.number_format, cell.border = f(bold=True), nf, box
mt.cell(row=tot + 1, column=1,
        value="Check vs Total Sales on KPI Definitions (should be 0):").font = f(size=9, italic=True)
mt.cell(row=tot + 1, column=4, value=f"=ROUND(D{tot}-'KPI Definitions'!$E$5,2)").font = f(size=9, italic=True)

# Seasonality grid: calendar month x year
head(mt, 4, 11, ["Month #", "Month", *years, "Average"])
mnames = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for i, name in enumerate(mnames):
    rr = 5 + i
    mt.cell(row=rr, column=11, value=i + 1)
    mt.cell(row=rr, column=12, value=name)
    for j, y in enumerate(years):
        col = 13 + j
        L = mt.cell(row=4, column=col).column_letter
        mt.cell(row=rr, column=col,
                value=f"=SUMIFS({D('R')},{D('W')},$K{rr},{D('V')},{L}$4)")
    mt.cell(row=rr, column=17, value=f"=AVERAGE(M{rr}:P{rr})")
    for c in range(11, 18):
        cell = mt.cell(row=rr, column=c)
        cell.font, cell.border = f(), box
        cell.number_format = "0" if c == 11 else ("@" if c == 12 else "$#,##0")
mt["K17"] = "Sales by calendar month and year. Compare the same month across years to see seasonality."
mt["K17"].font = f(size=9, italic=True, color="595959")

for col, w in zip("ABCDEFGHIJKLMNOPQ", [12, 8, 8, 12, 12, 12, 9, 12, 14, 3, 9, 8, 12, 12, 12, 12, 12]):
    mt.column_dimensions[col].width = w
mt.row_dimensions[4].height = 30
mt.freeze_panes = "A5"
mt.sheet_view.showGridLines = False

# Seasonality chart on this sheet
sea = LineChart()
sea.title = "Sales by calendar month (one line per year)"
sea.y_axis.number_format = "$#,##0"
sea.add_data(Reference(mt, min_col=13, max_col=16, min_row=4, max_row=16), titles_from_data=True)
sea.set_categories(Reference(mt, min_col=12, min_row=5, max_row=16))
style_axes(sea)
for s, colr in zip(sea.series, ["A6A6A6", "9DC3E6", "2E75B6", NAVY]):
    s.graphicalProperties.line.solidFill = colr
    s.graphicalProperties.line.width = 25400
    s.smooth = False
sea.height, sea.width = 8.5, 17
mt.add_chart(sea, "K19")

# ================================================================ Dashboard sheet
ds = wb.create_sheet("Dashboard", 0)
ds.sheet_view.showGridLines = False
ds.sheet_view.zoomScale = 90
ds.column_dimensions["A"].width = 2
for i in range(2, 16):
    ds.column_dimensions[ds.cell(row=1, column=i).column_letter].width = 12.5
ds.column_dimensions["P"].width = 2

ds.merge_cells("B1:O1")
ds["B1"] = "Superstore Sales Dashboard, 2015-2018"
ds["B1"].font = f(size=18, bold=True, color="FFFFFF")
ds["B1"].fill = PatternFill("solid", fgColor=NAVY)
ds["B1"].alignment = Alignment(vertical="center", indent=1)
ds.row_dimensions[1].height = 36
ds.merge_cells("B2:O2")
ds["B2"] = ("US retail order lines, gross of returns. All numbers are formulas on the Data sheet. "
            "Sales Growth compares 2018 with 2017.")
ds["B2"].font = f(size=9, italic=True, color="595959")

cards = [
    ("Total Sales", "E5", "$#,##0"),
    ("Total Profit", "E6", "$#,##0"),
    ("Profit Margin", "E7", "0.0%"),
    ("Total Orders", "E8", "#,##0"),
    ("Avg Order Value", "E10", "$#,##0.00"),
    ("Return Rate (orders)", "E12", "0.0%"),
    ("Sales Growth (YoY)", "E16", "+0.0%;-0.0%"),
]
ds.row_dimensions[4].height = 20
ds.row_dimensions[5].height = 38
for i, (label, ref, nf) in enumerate(cards):
    c1 = 2 + 2 * i
    L = ds.cell(row=4, column=c1).column_letter
    R = ds.cell(row=4, column=c1 + 1).column_letter
    ds.merge_cells(f"{L}4:{R}4")
    ds.merge_cells(f"{L}5:{R}5")
    lab = ds[f"{L}4"]
    lab.value = label
    lab.font = f(size=10, bold=True, color="FFFFFF")
    lab.fill = PatternFill("solid", fgColor=BLUE)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = ds[f"{L}5"]
    val.value = f"='KPI Definitions'!{ref}"
    val.font = f(size=20, bold=True, color=NAVY)
    val.fill = PatternFill("solid", fgColor=LIGHT)
    val.number_format = nf
    val.alignment = Alignment(horizontal="center", vertical="center")
    for cc in (c1, c1 + 1):
        ds.cell(row=4, column=cc).fill = PatternFill("solid", fgColor=BLUE)
        ds.cell(row=5, column=cc).fill = PatternFill("solid", fgColor=LIGHT)

def section(row, text):
    ds.merge_cells(f"B{row}:O{row}")
    c = ds[f"B{row}"]
    c.value = text
    c.font = f(size=11, bold=True, color=NAVY)
    c.fill = PatternFill("solid", fgColor=LIGHT)
    c.alignment = Alignment(vertical="center", indent=1)
    ds.row_dimensions[row].height = 18


section(6, "Sales trend")

# Chart 1: monthly trend (full width)
tr = LineChart()
tr.title = "Monthly Sales and Profit"
tr.y_axis.number_format = "$#,##0"
tr.y_axis.title = "USD"
tr.add_data(Reference(mt, min_col=4, max_col=5, min_row=4, max_row=lastm), titles_from_data=True)
tr.set_categories(Reference(mt, min_col=1, min_row=5, max_row=lastm))
tr.x_axis.tickLblSkip = 3
tr.x_axis.tickMarkSkip = 3
style_axes(tr)
for s, colr in zip(tr.series, [BLUE, "ED7D31"]):
    s.graphicalProperties.line.solidFill = colr
    s.graphicalProperties.line.width = 28575
    s.smooth = False
tr.height, tr.width = 8.5, 34.5
ds.add_chart(tr, "B7")


def bars(title, hdr_row, first_row, last_row, anchor):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.y_axis.number_format = "$#,##0"
    ch.add_data(Reference(kpi, min_col=3, max_col=4, min_row=hdr_row, max_row=last_row),
                titles_from_data=True)
    ch.set_categories(Reference(kpi, min_col=2, min_row=first_row, max_row=last_row))
    style_axes(ch)
    for s, colr in zip(ch.series, [BLUE, "ED7D31"]):
        s.graphicalProperties.solidFill = colr
    ch.gapWidth = 80
    ch.height, ch.width = 8, 17
    ds.add_chart(ch, anchor)


bars("Sales and Profit by Category", 28, 29, 31, "B25")
bars("Sales and Profit by Region", 35, 36, 39, "I25")

section(24, "Category and region")
section(42, "Sub-category performance (ranked by profit)")
section(64, "Products: top 10 and bottom 10 by profit")
section(85, "Discounting")


def hbar(title, sh, cat_col, val_col, hdr, first, last, anchor, colour, nf, height=10):
    ch = BarChart()
    ch.type = "bar"
    ch.title = title
    ch.add_data(Reference(sh, min_col=val_col, min_row=hdr, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(sh, min_col=cat_col, min_row=first, max_row=last))
    style_axes(ch)
    ch.legend = None
    ch.x_axis.scaling.orientation = "maxMin"  # rank 1 at the top
    ch.y_axis.crosses = "max"                 # keep the value axis at the bottom
    ch.y_axis.number_format = nf
    ch.series[0].graphicalProperties.solidFill = colour
    ch.series[0].invertIfNegative = False
    ch.gapWidth = 40
    ch.height, ch.width = height, 17
    ds.add_chart(ch, anchor)


hbar("Profit by Sub-Category", pa, 12, 14, 5, 6, 22, "B43", BLUE, "$#,##0", 10)
hbar("Profit Margin by Sub-Category", pa, 12, 15, 5, 6, 22, "I43", "70AD47", "0%", 10)
hbar("Top 10 Products by Profit", pa, 2, 5, 28, 29, 38, "B65", BLUE, "$#,##0", 9.5)
hbar("Bottom 10 Products by Profit", pa, 2, 5, 42, 43, 52, "I65", "C00000", "$#,##0", 9.5)


def dcols(title, val_col, anchor, colour, nf):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.add_data(Reference(da, min_col=val_col, min_row=12, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(da, min_col=1, min_row=13, max_row=16))
    style_axes(ch)
    ch.legend = None
    ch.y_axis.number_format = nf
    ch.series[0].graphicalProperties.solidFill = colour
    ch.series[0].invertIfNegative = False
    ch.x_axis.title = "Discount band"
    ch.gapWidth = 60
    ch.height, ch.width = 8, 17
    ds.add_chart(ch, anchor)


dcols("Profit Margin by Discount Band", 7, "B86", BLUE, "0%")
dcols("Profit by Discount Band", 5, "I86", "ED7D31", "$#,##0")

ds["B104"] = ("Filters: Year / Region / Segment slicers need a PivotTable. See the README for the "
              "steps. Details and checks are on the Monthly Trend, Product Analysis and Discount "
              "Analysis sheets.")
ds["B104"].font = f(size=9, italic=True, color="595959")

# print setup: landscape, one page wide
ds.page_setup.orientation = "landscape"
ds.page_setup.fitToWidth = 1
ds.page_setup.fitToHeight = 0
ds.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ds.print_area = "A1:P105"

# keep the data sheets out of the way in the preview render only
if PREVIEW:
    for name in ("KPI Definitions", "Cleaning Log", "Data", "Returns", "Regional Managers",
                 "Product Analysis", "Discount Analysis", "Product Summary"):
        wb[name].sheet_state = "hidden"

wb.active = 0
wb.save(OUT)
print("saved", OUT, "| months", first, lastm)
