"""
Add the analysis tables to superstore_beginner.xlsx.

Run after build_workbook.py and before add_dashboard.py:
    python scripts/add_analysis.py

Adds
  * Data!AD 'Discount Band'   (formula; band limits are inputs on Discount Analysis)
  * Product Summary           one row per Product Name (SUMIFS on Data)
  * Product Analysis          sub-category table, sub-categories ranked by profit,
                              top 10 and bottom 10 products by profit
  * Discount Analysis         results by discount band, by exact rate, by category,
                              by calendar month, and discount-vs-margin correlation

Charts are added by add_dashboard.py (openpyxl drops charts when it re-loads a file).
Everything is a formula except labels and the blue input cells.
"""
from pathlib import Path
from copy import copy
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "superstore_beginner.xlsx"
orders = pd.read_csv(ROOT / "data/clean/superstore_clean.csv")

wb = load_workbook(SRC)
for name in ("Product Analysis", "Product Summary", "Discount Analysis"):
    if name in wb.sheetnames:
        del wb[name]
data = wb["Data"]
last = data.max_row
assert data["AC1"].value == "First Line of Order", "Data sheet layout changed"

FONT = "Arial"
f = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)  # noqa: E731
NAVY, GREY = "1F3864", "F2F2F2"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = f(color="0000FF")  # blue = editable input
D = lambda col: f"Data!${col}$2:${col}${last}"  # noqa: E731
DA = "'Discount Analysis'"


def head(ws, row, col, labels, height=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if height:
        ws.row_dimensions[row].height = height


def put(ws, row, col, value, nf=None, bold=False, border=True, align=None, font=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or f(bold=bold)
    if nf:
        c.number_format = nf
    if border:
        c.border = box
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = f(size=14, bold=True)
    ws["A2"] = note
    ws["A2"].font = f(size=9, italic=True, color="595959")
    ws.sheet_view.showGridLines = False


# ====================================================== Discount Analysis sheet
da = wb.create_sheet("Discount Analysis", 0)
title(da, "Discount analysis: how discounting relates to profit",
      "Gross figures. Blue cells are inputs: change the band limits and everything, "
      "including the Data sheet's Discount Band column, updates.")

da["A4"] = "Band settings"
da["A4"].font = f(bold=True)
head(da, 5, 1, ["Band label", "Upper limit (included)", "Meaning"], 30)
bands = [("0%", 0, "No discount"),
         ("1-20%", 0.2, "Light discount (0.2 is the most common rate)"),
         ("21-40%", 0.4, "Heavy discount"),
         ("41%+", None, "Deep discount (everything above the previous limit)")]
for i, (lab, up, meaning) in enumerate(bands):
    r = 6 + i
    put(da, r, 1, lab, font=INPUT)
    put(da, r, 2, up, "0%", font=INPUT)
    put(da, r, 3, meaning)
    if up is None:
        da.cell(row=r, column=2).value = "no limit"
        da.cell(row=r, column=2).font = f()
        da.cell(row=r, column=2).alignment = Alignment(horizontal="right")

# Discount Band column on the Data sheet
col = 30  # AD
hdr = data.cell(row=1, column=col, value="Discount Band")
src_h = data["AC1"]
hdr.font, hdr.fill, hdr.border, hdr.alignment = (copy(src_h.font), copy(src_h.fill),
                                                 copy(src_h.border), copy(src_h.alignment))
for r in range(2, last + 1):
    c = data.cell(row=r, column=col,
                  value=(f"=IF(T{r}<={DA}!$B$6,{DA}!$A$6,IF(T{r}<={DA}!$B$7,{DA}!$A$7,"
                         f"IF(T{r}<={DA}!$B$8,{DA}!$A$8,{DA}!$A$9)))"))
    c.font, c.number_format = f(), "@"
data.column_dimensions["AD"].width = 13
data.auto_filter.ref = f"A1:AD{last}"

# Table 1: by band
da["A11"] = "1. Results by discount band"
da["A11"].font = f(bold=True)
head(da, 12, 1, ["Band", "Lines", "% of Lines", "Sales", "Profit", "% of Total Profit",
                 "Profit Margin", "Avg Discount", "Loss-making Lines %"], 32)
for i in range(4):
    r = 13 + i
    put(da, r, 1, f"=A{6+i}")
    put(da, r, 2, f"=COUNTIFS({D('AD')},$A{r})", "#,##0")
    put(da, r, 3, f"=IF($B$17=0,0,B{r}/$B$17)", "0.0%")
    put(da, r, 4, f"=SUMIFS({D('R')},{D('AD')},$A{r})", "$#,##0")
    put(da, r, 5, f"=SUMIFS({D('U')},{D('AD')},$A{r})", "$#,##0;-$#,##0")
    put(da, r, 6, f"=IF($E$17=0,0,E{r}/$E$17)", "0.0%")
    put(da, r, 7, f"=IF(D{r}=0,0,E{r}/D{r})", "0.0%")
    put(da, r, 8, f"=IFERROR(AVERAGEIFS({D('T')},{D('AD')},$A{r}),0)", "0.0%")
    put(da, r, 9, f'=IF(B{r}=0,0,COUNTIFS({D("AD")},$A{r},{D("U")},"<0")/B{r})', "0.0%")
put(da, 17, 1, "Total", bold=True)
put(da, 17, 2, "=SUM(B13:B16)", "#,##0", True)
put(da, 17, 3, "=SUM(C13:C16)", "0.0%", True)
put(da, 17, 4, "=SUM(D13:D16)", "$#,##0", True)
put(da, 17, 5, "=SUM(E13:E16)", "$#,##0;-$#,##0", True)
put(da, 17, 6, "=SUM(F13:F16)", "0.0%", True)
put(da, 17, 7, "=IF(D17=0,0,E17/D17)", "0.0%", True)
put(da, 17, 8, f"=AVERAGE({D('T')})", "0.0%", True)
put(da, 17, 9, f'=COUNTIF({D("U")},"<0")/COUNT({D("U")})', "0.0%", True)
put(da, 18, 1, "Lines above the 1-20% band", border=False, font=f(size=9, italic=True))
put(da, 18, 5, "=E15+E16", "$#,##0;-$#,##0", border=False, font=f(size=9, italic=True))
put(da, 18, 6, "<- combined profit of the two heaviest discount bands",
    border=False, font=f(size=9, italic=True))

# Table 2: by exact rate
da["A21"] = "2. Results by exact discount rate"
da["A21"].font = f(bold=True)
head(da, 22, 1, ["Discount", "Lines", "Sales", "Profit", "Profit Margin"], 20)
rates = sorted(orders["Discount"].unique())
for i, v in enumerate(rates):
    r = 23 + i
    put(da, r, 1, float(v), "0%", font=INPUT)
    put(da, r, 2, f"=COUNTIFS({D('T')},$A{r})", "#,##0")
    put(da, r, 3, f"=SUMIFS({D('R')},{D('T')},$A{r})", "$#,##0")
    put(da, r, 4, f"=SUMIFS({D('U')},{D('T')},$A{r})", "$#,##0;-$#,##0")
    put(da, r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", "0.0%")
r_rates_end = 23 + len(rates) - 1
put(da, r_rates_end + 1, 1, "Total", bold=True)
put(da, r_rates_end + 1, 2, f"=SUM(B23:B{r_rates_end})", "#,##0", True)
put(da, r_rates_end + 1, 3, f"=SUM(C23:C{r_rates_end})", "$#,##0", True)
put(da, r_rates_end + 1, 4, f"=SUM(D23:D{r_rates_end})", "$#,##0;-$#,##0", True)
put(da, r_rates_end + 1, 5, f"=IF(C{r_rates_end+1}=0,0,D{r_rates_end+1}/C{r_rates_end+1})", "0.0%", True)

# Table 3: band x category margin
t3 = r_rates_end + 4
da.cell(row=t3, column=1, value="3. Profit margin by discount band and category").font = f(bold=True)
cats = sorted(orders["Category"].unique())
head(da, t3 + 1, 1, ["Band", *cats], 20)
for i in range(4):
    r = t3 + 2 + i
    put(da, r, 1, f"=A{6+i}")
    for j in range(len(cats)):
        L = get_column_letter(2 + j)
        crit = f"{D('AD')},$A{r},{D('O')},{L}${t3+1}"
        put(da, r, 2 + j,
            f'=IFERROR(SUMIFS({D("U")},{crit})/SUMIFS({D("R")},{crit}),"n/a")', "0.0%", align="right")
da.cell(row=t3 + 6, column=1, value="n/a = no order lines in that combination.").font = f(size=9, italic=True)
da.conditional_formatting.add(f"B{t3+2}:D{t3+5}",
                              CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, color="C00000")))

# Table 4: calendar month
t4 = t3 + 9
da.cell(row=t4, column=1, value="4. Is peak season driven by deeper discounts? (by calendar month)").font = f(bold=True)
head(da, t4 + 1, 1, ["Month #", "Month", "Sales", "Avg Discount", "Profit Margin",
                     f'="Share of Sales discounted above "&TEXT($B$7,"0%")'], 44)
mnames = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for i, mn in enumerate(mnames):
    r = t4 + 2 + i
    put(da, r, 1, i + 1, "0")
    put(da, r, 2, mn)
    put(da, r, 3, f"=SUMIFS({D('R')},{D('W')},$A{r})", "$#,##0")
    put(da, r, 4, f"=AVERAGEIFS({D('T')},{D('W')},$A{r})", "0.0%")
    put(da, r, 5, f"=IF(C{r}=0,0,SUMIFS({D('U')},{D('W')},$A{r})/C{r})", "0.0%")
    put(da, r, 6, f'=IF(C{r}=0,0,SUMIFS({D("R")},{D("W")},$A{r},{D("T")},">"&$B$7)/C{r})', "0.0%")
m1, m2 = t4 + 2, t4 + 13
c1 = t4 + 15
put(da, c1, 1, "Correlation across the 12 calendar months, Avg Discount vs Profit Margin",
    border=False, font=f(bold=True))
put(da, c1, 6, f"=CORREL(D{m1}:D{m2},E{m1}:E{m2})", "0.00", True)
put(da, c1 + 1, 1, "Correlation across all 48 individual months, Avg Discount vs Profit Margin",
    border=False, font=f(bold=True))
put(da, c1 + 1, 6, "=CORREL(L13:L60,M13:M60)", "0.00", True)
put(da, c1 + 2, 1, ("Reading it: 0 = no relationship, -1 = higher discount always means lower margin. "
                    "Correlation shows association, not cause."),
    border=False, font=f(size=9, italic=True, color="595959"))

# Helper: 48 months (Year-Month, avg discount, margin) for the correlation
da["K11"] = "Helper: monthly values for the 48-month correlation"
da["K11"].font = f(bold=True)
head(da, 12, 11, ["Year-Month", "Avg Discount", "Profit Margin"], 32)
for i in range(48):
    r = 13 + i
    put(da, r, 11, f"='Monthly Trend'!A{5+i}", "@")
    put(da, r, 12, f"=IFERROR(AVERAGEIFS({D('T')},{D('X')},K{r}),0)", "0.0%")
    put(da, r, 13, f"='Monthly Trend'!F{5+i}", "0.0%")

for colL, w in zip("ABCDEFGHIJKLM", [24, 22, 13, 13, 13, 17, 13, 13, 14, 3, 12, 13, 13]):
    da.column_dimensions[colL].width = w
da.row_dimensions[t4 + 1].height = 44

# ======================================================== Product Summary sheet
ps = wb.create_sheet("Product Summary")
title(ps, "Product summary (one row per Product Name)",
      "Product Name is used, not Product ID, because 32 Product IDs are shared by two products. "
      "Match Key escapes the characters * ? ~ so SUMIFS treats them as plain text. "
      "Rank Key = Profit plus a tiny row-based amount so ties still rank in a fixed order.")
head(ps, 4, 1, ["Product Name", "Match Key", "Sales", "Profit", "Profit Margin", "Rank Key"], 28)
names = sorted(orders["Product Name"].unique())
p1, pn = 5, 5 + len(names) - 1
for i, nm in enumerate(names):
    r = p1 + i
    put(ps, r, 1, nm, border=False)
    put(ps, r, 2, f'=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(A{r},"~","~~"),"*","~*"),"?","~?")', border=False)
    put(ps, r, 3, f"=SUMIFS({D('R')},{D('Q')},$B{r})", "$#,##0.00", border=False)
    put(ps, r, 4, f"=SUMIFS({D('U')},{D('Q')},$B{r})", "$#,##0.00;-$#,##0.00", border=False)
    put(ps, r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", "0.0%", border=False)
    put(ps, r, 6, f"=D{r}+ROW()*0.0000001", "0.0000000", border=False)
put(ps, pn + 2, 1, "Total (should equal Total Sales / Total Profit)", bold=True, border=False)
put(ps, pn + 2, 3, f"=SUM(C{p1}:C{pn})", "$#,##0.00", True, False)
put(ps, pn + 2, 4, f"=SUM(D{p1}:D{pn})", "$#,##0.00;-$#,##0.00", True, False)
put(ps, pn + 3, 1, "Difference vs KPI Definitions", border=False, font=f(size=9, italic=True))
put(ps, pn + 3, 3, f"=ROUND(C{pn+2}-'KPI Definitions'!$E$5,2)", "0.00", border=False, font=f(size=9, italic=True))
put(ps, pn + 3, 4, f"=ROUND(D{pn+2}-'KPI Definitions'!$E$6,2)", "0.00", border=False, font=f(size=9, italic=True))
for colL, w in zip("ABCDEF", [55, 30, 13, 13, 13, 16]):
    ps.column_dimensions[colL].width = w
ps.freeze_panes = "A5"
ps.auto_filter.ref = f"A4:F{pn}"
wb.move_sheet("Product Summary", offset=wb.sheetnames.index("Data") + 1 - wb.sheetnames.index("Product Summary"))

# ======================================================= Product Analysis sheet
pa = wb.create_sheet("Product Analysis", 0)
title(pa, "Sub-category and product analysis",
      "Gross figures. Sections: (1) sub-category table, (2) sub-categories ranked by profit, "
      "(3) top 10 and (4) bottom 10 products by profit.")

subs = (orders.groupby("Sub-Category")["Category"].first().sort_index())
n = len(subs)
s1, sn = 6, 5 + n  # 6..22
pa["A4"] = "1. Sub-category summary"
pa["A4"].font = f(bold=True)
head(pa, 5, 1, ["Sub-Category", "Category", "Sales", "Profit", "Profit Margin",
                "Avg Discount", "Units", "Rank Key"], 30)
for i, sc in enumerate(subs.index):
    r = s1 + i
    put(pa, r, 1, sc)
    put(pa, r, 2, f"=INDEX({D('O')},MATCH($A{r},{D('P')},0))")
    put(pa, r, 3, f"=SUMIFS({D('R')},{D('P')},$A{r})", "$#,##0")
    put(pa, r, 4, f"=SUMIFS({D('U')},{D('P')},$A{r})", "$#,##0;-$#,##0")
    put(pa, r, 5, f"=IF(C{r}=0,0,D{r}/C{r})", "0.0%")
    put(pa, r, 6, f"=AVERAGEIFS({D('T')},{D('P')},$A{r})", "0.0%")
    put(pa, r, 7, f"=SUMIFS({D('S')},{D('P')},$A{r})", "#,##0")
    put(pa, r, 8, f"=D{r}+ROW()*0.0000001", "0.0000000")
tr = sn + 1
put(pa, tr, 1, "Total", bold=True)
put(pa, tr, 2, "", bold=True)
put(pa, tr, 3, f"=SUM(C{s1}:C{sn})", "$#,##0", True)
put(pa, tr, 4, f"=SUM(D{s1}:D{sn})", "$#,##0;-$#,##0", True)
put(pa, tr, 5, f"=IF(C{tr}=0,0,D{tr}/C{tr})", "0.0%", True)
put(pa, tr, 6, f"=AVERAGE({D('T')})", "0.0%", True)
put(pa, tr, 7, f"=SUM(G{s1}:G{sn})", "#,##0", True)
put(pa, tr + 1, 1, "Difference vs Total Sales", border=False, font=f(size=9, italic=True))
put(pa, tr + 1, 3, f"=ROUND(C{tr}-'KPI Definitions'!$E$5,2)", "0.00", border=False, font=f(size=9, italic=True))

# ranked sub-categories (K..Q)
pa["K4"] = "2. Sub-categories ranked by profit (highest first)"
pa["K4"].font = f(bold=True)
head(pa, 5, 11, ["Rank", "Sub-Category", "Category", "Profit", "Profit Margin", "Avg Discount", "Row"], 30)
for i in range(n):
    r = s1 + i
    put(pa, r, 11, i + 1, "0")
    put(pa, r, 17, f"=MATCH(LARGE($H${s1}:$H${sn},K{r}),$H${s1}:$H${sn},0)", "0")
    put(pa, r, 12, f"=INDEX($A${s1}:$A${sn},$Q{r})")
    put(pa, r, 13, f"=INDEX($B${s1}:$B${sn},$Q{r})")
    put(pa, r, 14, f"=INDEX($D${s1}:$D${sn},$Q{r})", "$#,##0;-$#,##0")
    put(pa, r, 15, f"=INDEX($E${s1}:$E${sn},$Q{r})", "0.0%")
    put(pa, r, 16, f"=INDEX($F${s1}:$F${sn},$Q{r})", "0.0%")
pa.conditional_formatting.add(f"N{s1}:O{sn}",
                              CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, color="C00000")))
pa.conditional_formatting.add(f"D{s1}:E{sn}",
                              CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, color="C00000")))


def product_block(top, label, fn, hdr_row):
    pa.cell(row=top, column=1, value=label).font = f(bold=True)
    head(pa, hdr_row, 1, ["Rank", "Product (short)", "Product Name", "Sales", "Profit",
                          "Profit Margin", "Row"], 30)
    for i in range(10):
        r = hdr_row + 1 + i
        rng = f"'Product Summary'!$F${p1}:$F${pn}"
        put(pa, r, 1, i + 1, "0")
        put(pa, r, 7, f"=MATCH({fn}({rng},A{r}),{rng},0)", "0")
        put(pa, r, 3, f"=INDEX('Product Summary'!$A${p1}:$A${pn},$G{r})")
        put(pa, r, 2, f'=LEFT(C{r},34)&IF(LEN(C{r})>34,"...","")')
        put(pa, r, 4, f"=INDEX('Product Summary'!$C${p1}:$C${pn},$G{r})", "$#,##0")
        put(pa, r, 5, f"=INDEX('Product Summary'!$D${p1}:$D${pn},$G{r})", "$#,##0;-$#,##0")
        put(pa, r, 6, f"=INDEX('Product Summary'!$E${p1}:$E${pn},$G{r})", "0.0%")
    pa.conditional_formatting.add(f"E{hdr_row+1}:F{hdr_row+10}",
                                  CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, color="C00000")))


TOP_HDR, BOT_HDR = 28, 42
product_block(27, "3. Top 10 products by profit", "LARGE", TOP_HDR)
product_block(41, "4. Bottom 10 products by profit", "SMALL", BOT_HDR)

for colL, w in zip("ABCDEFGHIJKLMNOPQ", [20, 36, 52, 13, 13, 13, 9, 14, 3, 3, 7, 18, 16, 13, 13, 13, 7]):
    pa.column_dimensions[colL].width = w
pa.freeze_panes = "A4"

wb.save(SRC)
print("saved", SRC, "| products", len(names), "| sub-categories", n)
