"""
Build superstore_beginner.xlsx from the cleaned CSVs.

Sheets: KPI Definitions | Cleaning Log | Data | Returns | Regional Managers
Every KPI and every derived column is an Excel formula, so the numbers
recalculate if the data changes.

Run from the 01-beginner folder:  python scripts/build_workbook.py
Then recalculate (LibreOffice) or just open in Excel.
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "superstore_beginner.xlsx"

orders = pd.read_csv(ROOT / "data/clean/superstore_clean.csv", dtype={"Postal Code": str},
                     parse_dates=["Order Date", "Ship Date"])
returns = pd.read_csv(ROOT / "data/clean/returns.csv")
managers = pd.read_csv(ROOT / "data/clean/regional_managers.csv")
log = pd.read_csv(ROOT / "docs/cleaning_log.csv")

FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True)
f_note = Font(name=FONT, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_derived = PatternFill("solid", fgColor="2E75B6")
fill_band = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, labels, fill=fill_head):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.border = f_head, fill, box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


# --------------------------------------------------------------------------- Data
ws_data = wb.active
ws_data.title = "Data"
base_cols = list(orders.columns)  # 21 columns, A..U
derived = [
    ("Order Year", "=YEAR(C{r})", "0"),
    ("Order Month", "=MONTH(C{r})", "0"),
    ("Year-Month", '=YEAR(C{r})&"-"&TEXT(MONTH(C{r}),"00")', "@"),
    ("Quarter", '="Q"&ROUNDUP(MONTH(C{r})/3,0)', "@"),
    ("Ship Days", "=D{r}-C{r}", "0"),
    ("Line Profit Margin", "=IF(R{r}=0,0,U{r}/R{r})", "0.0%"),
    ("Is Returned", "=IF(COUNTIF(Returns!$A$2:$A${ret_last},B{r})>0,1,0)", "0"),
    ("First Line of Order", "=IF(B{r}<>B{p},1,0)", "0"),
]
header(ws_data, 1, base_cols)
for j, (name, _, _) in enumerate(derived, len(base_cols) + 1):
    c = ws_data.cell(row=1, column=j, value=name)
    c.font, c.fill, c.border = f_head, fill_derived, box
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

last = len(orders) + 1  # last data row on the Data sheet
ret_last = len(returns) + 1
fmt = {"Order Date": "yyyy-mm-dd", "Ship Date": "yyyy-mm-dd", "Sales": "#,##0.00",
       "Profit": "#,##0.00", "Discount": "0%", "Postal Code": "@"}
for i, row in enumerate(orders.itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        if isinstance(v, pd.Timestamp):
            v = v.to_pydatetime()
        c = ws_data.cell(row=i, column=j, value=v)
        c.font = f_norm
        if base_cols[j - 1] in fmt:
            c.number_format = fmt[base_cols[j - 1]]
    for j, (_, formula, nf) in enumerate(derived, len(base_cols) + 1):
        c = ws_data.cell(row=i, column=j,
                         value=formula.format(r=i, p=i - 1, ret_last=ret_last))
        c.font, c.number_format = f_norm, nf
# first data row has no previous row to compare to
ws_data.cell(row=2, column=len(base_cols) + len(derived), value=1)

widths = {"A": 8, "B": 16, "C": 12, "D": 12, "E": 15, "F": 12, "G": 20, "H": 12, "I": 14,
          "J": 16, "K": 16, "L": 11, "M": 10, "N": 17, "O": 16, "P": 13, "Q": 45, "R": 11,
          "S": 9, "T": 9, "U": 11}
for k, w in widths.items():
    ws_data.column_dimensions[k].width = w
for j in range(len(base_cols) + 1, len(base_cols) + len(derived) + 1):
    ws_data.column_dimensions[get_column_letter(j)].width = 13
ws_data.row_dimensions[1].height = 32
ws_data.freeze_panes = "C2"
ws_data.auto_filter.ref = f"A1:{get_column_letter(len(base_cols) + len(derived))}{last}"

# ------------------------------------------------------------------------ Returns
ws_ret = wb.create_sheet("Returns")
header(ws_ret, 1, ["Order ID", "Returned"])
for i, r in enumerate(returns.itertuples(index=False), 2):
    ws_ret.cell(row=i, column=1, value=r[0]).font = f_norm
    ws_ret.cell(row=i, column=2, value=r[1]).font = f_norm
ws_ret.column_dimensions["A"].width = 18
ws_ret.column_dimensions["B"].width = 10
ws_ret.freeze_panes = "A2"

ws_mgr = wb.create_sheet("Regional Managers")
header(ws_mgr, 1, ["Regional Manager", "Region"])
for i, r in enumerate(managers.itertuples(index=False), 2):
    ws_mgr.cell(row=i, column=1, value=r[0]).font = f_norm
    ws_mgr.cell(row=i, column=2, value=r[1]).font = f_norm
ws_mgr.column_dimensions["A"].width = 22
ws_mgr.column_dimensions["B"].width = 12

# ------------------------------------------------------------------- Cleaning Log
ws_log = wb.create_sheet("Cleaning Log")
ws_log["A1"] = "Cleaning log: every change made to the raw file"
ws_log["A1"].font = f_title
ws_log["A2"] = ("Source: superstore_raw.csv (see README). Reproduce with "
                "scripts/clean_superstore.py.")
ws_log["A2"].font = f_note
header(ws_log, 4, list(log.columns))
for i, r in enumerate(log.itertuples(index=False), 5):
    for j, v in enumerate(r, 1):
        c = ws_log.cell(row=i, column=j, value=v)
        c.font, c.alignment, c.border = f_norm, wrap, box
        if j in (4, 5):
            c.number_format = "#,##0"
            c.alignment = Alignment(vertical="top", horizontal="right")
for col, w in zip("ABCDE", [6, 40, 60, 60, 14]):
    ws_log.column_dimensions[col].width = w
ws_log.column_dimensions["F"].width = 16
ws_log.freeze_panes = "A5"

# --------------------------------------------------------------- KPI Definitions
ws = wb.create_sheet("KPI Definitions", 0)
ws["A1"] = "Superstore: KPI definitions"
ws["A1"].font = f_title
ws["A2"] = ("Values are live formulas on the Data sheet. Sales, Profit and Quantity are summed "
            "across all order lines. Returned orders are included in headline Sales and Profit "
            "(gross view); KPI 9 shows how much of that was later returned.")
ws["A2"].font = f_note
header(ws, 4, ["#", "KPI", "Definition", "Calculation", "Value", "Dashboard use", "Why it matters"])

D = lambda col: f"Data!${col}$2:${col}${last}"  # noqa: E731
kpis = [
    ("Total Sales", "Revenue from all order lines, before returns.",
     "SUM of Sales", f"=SUM({D('R')})", "$#,##0", "Headline card",
     "Size of the business; the base for every ratio below."),
    ("Total Profit", "Profit from all order lines, after discounts.",
     "SUM of Profit", f"=SUM({D('U')})", "$#,##0", "Headline card",
     "Shows whether the sales are actually making money."),
    ("Profit Margin", "Share of each sales dollar kept as profit.",
     "Total Profit / Total Sales", "=IF(E5=0,0,E6/E5)", "0.0%", "Headline card",
     "Compares categories and regions fairly regardless of size."),
    ("Total Orders", "Number of distinct orders (an order can have several lines).",
     "Count of lines flagged 'First Line of Order'", f"=SUM({D('AC')})", "#,##0",
     "Headline card", "Counting rows would overstate orders; this counts each Order ID once."),
    ("Units Sold", "Total quantity of items sold.", "SUM of Quantity",
     f"=SUM({D('S')})", "#,##0", "Supporting", "Volume view, separate from price effects."),
    ("Average Order Value", "Average sales per order.", "Total Sales / Total Orders",
     "=IF(E8=0,0,E5/E8)", "$#,##0.00", "Headline card",
     "Tracks basket size; lifts with bundling or minimum-spend offers."),
    ("Average Discount", "Mean discount rate across order lines.", "AVERAGE of Discount",
     f"=AVERAGE({D('T')})", "0.0%", "Supporting",
     "Discount depth is the usual driver of low margins."),
    ("Return Rate (orders)", "Share of orders that were returned.",
     "Returned orders / Total Orders",
     f"=IF(E8=0,0,SUMPRODUCT({D('AB')},{D('AC')})/E8)", "0.0%", "Headline card",
     "Returns erase revenue and add handling cost."),
    ("Returned Sales Share", "Share of Sales that sits on returned orders.",
     "Sales where Is Returned = 1 / Total Sales",
     f"=IF(E5=0,0,SUMIFS({D('R')},{D('AB')},1)/E5)", "0.0%", "Supporting",
     "Shows the revenue at risk from returns, not just the order count."),
    ("Average Ship Days", "Average days from order to shipment.", "AVERAGE of Ship Days",
     f"=AVERAGE({D('Z')})", "0.0", "Supporting", "Delivery speed; compare across Ship Mode."),
    ("Loss-making Lines", "Share of order lines with negative profit.",
     "Lines with Profit < 0 / all lines",
     f'=COUNTIF({D("U")},"<0")/COUNT({D("U")})', "0.0%", "Supporting",
     "Finds how widespread losses are before looking at causes."),
    ("Sales Growth (latest year vs prior)", "Change in yearly Sales, last year vs the year before.",
     "(Sales latest year / Sales prior year) - 1", "=IF(C28=0,0,C29/C28-1)", "0.0%",
     "Headline card", "The simplest trend signal for an executive reader."),
]
# Place growth formula after we know the year table position (set below).
for i, (kpi, definition, calc, formula, nf, use, why) in enumerate(kpis, 1):
    r = 4 + i
    vals = [i, kpi, definition, calc, formula, use, why]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.alignment, c.border = f_norm, wrap, box
        if i % 2 == 0:
            c.fill = fill_band
    ws.cell(row=r, column=5).number_format = nf
    ws.cell(row=r, column=5).font = f_bold
    ws.cell(row=r, column=5).alignment = Alignment(vertical="top", horizontal="right")
    ws.cell(row=r, column=2).font = f_bold

# --- supporting tables
years = sorted(orders["Order Date"].dt.year.unique())
cats = sorted(orders["Category"].unique())
regs = sorted(orders["Region"].unique())

t = 4 + len(kpis) + 3  # title row of first supporting table
ws.cell(row=t - 1, column=1, value="Supporting tables (sanity checks and dashboard source)").font = f_title


def table(top, title, key_label, keys, key_col, extra_growth=False):
    ws.cell(row=top, column=2, value=title).font = f_bold
    heads = [key_label, "Sales", "Profit", "Profit Margin"]
    if extra_growth:
        heads.append("Sales Growth")
    for j, h in enumerate(heads, 2):
        c = ws.cell(row=top + 1, column=j, value=h)
        c.font, c.fill, c.border = f_head, fill_head, box
        c.alignment = Alignment(horizontal="center")
    for i, k in enumerate(keys):
        r = top + 2 + i
        ws.cell(row=r, column=2, value=k)
        ws.cell(row=r, column=3, value=f"=SUMIFS({D('R')},{D(key_col)},$B{r})")
        ws.cell(row=r, column=4, value=f"=SUMIFS({D('U')},{D(key_col)},$B{r})")
        ws.cell(row=r, column=5, value=f"=IF(C{r}=0,0,D{r}/C{r})")
        if extra_growth and i > 0:
            ws.cell(row=r, column=6, value=f"=IF(C{r-1}=0,0,C{r}/C{r-1}-1)")
        for j in range(2, 2 + len(heads)):
            c = ws.cell(row=r, column=j)
            c.font, c.border = f_norm, box
            c.number_format = {2: "0" if key_col == "V" else "@", 3: "$#,##0", 4: "$#,##0",
                               5: "0.0%", 6: "0.0%"}[j]
            c.alignment = Alignment(horizontal="left" if j == 2 else "right")
    total = top + 2 + len(keys)
    ws.cell(row=total, column=2, value="Total").font = f_bold
    ws.cell(row=total, column=3, value=f"=SUM(C{top+2}:C{total-1})")
    ws.cell(row=total, column=4, value=f"=SUM(D{top+2}:D{total-1})")
    ws.cell(row=total, column=5, value=f"=IF(C{total}=0,0,D{total}/C{total})")
    for j, nf in zip((3, 4, 5), ("$#,##0", "$#,##0", "0.0%")):
        c = ws.cell(row=total, column=j)
        c.font, c.number_format, c.border = f_bold, nf, box
    ws.cell(row=total, column=2).border = box
    return top + 2, total


y_first, y_total = table(t, "By Order Year", "Year", years, "V", extra_growth=True)
c_first, c_total = table(y_total + 2, "By Category", "Category", cats, "O")
g_first, g_total = table(c_total + 2, "By Region", "Region", regs, "M")

# Point the growth KPI at the year table (latest year row vs the one before)
latest, prior = y_total - 1, y_total - 2
ws["E16"] = f"=IF(C{prior}=0,0,C{latest}/C{prior}-1)"
ws["D16"] = f"(Sales {years[-1]} / Sales {years[-2]}) - 1"

ws.cell(row=g_total + 2, column=2,
        value=("Check: the three 'Total' rows above should all equal KPI 1 and KPI 2. "
               "Difference vs Total Sales:")).font = f_note
ws.cell(row=g_total + 3, column=2, value="Year / Category / Region").font = f_note
for j, tot in zip((3, 4, 5), (y_total, c_total, g_total)):
    c = ws.cell(row=g_total + 3, column=j, value=f"=ROUND(C{tot}-$E$5,2)")
    c.font, c.number_format = f_note, "0.00"

for col, w in zip("ABCDEFG", [5, 30, 42, 38, 16, 16, 50]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
ws.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT, "| data rows", last - 1, "| year table rows", y_first, y_total)
