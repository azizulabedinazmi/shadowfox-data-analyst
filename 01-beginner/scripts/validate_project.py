"""Validate the cleaned data and workbook structure.

Run from the 01-beginner folder after rebuilding the project:
    python scripts/validate_project.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
orders = pd.read_csv(ROOT / "data/clean/superstore_clean.csv")
returns = pd.read_csv(ROOT / "data/clean/returns.csv")
managers = pd.read_csv(ROOT / "data/clean/regional_managers.csv")
workbook_path = ROOT / "superstore_beginner.xlsx"

expected_sheets = {
    "Dashboard",
    "Monthly Trend",
    "Product Analysis",
    "Product Summary",
    "Discount Analysis",
    "KPI Definitions",
    "Data",
    "Returns",
    "Regional Managers",
    "Cleaning Log",
}
assert len(orders) == 9_993
assert len(returns) == 296
assert len(managers) == 4
assert orders["Order ID"].nunique() == 5_009
assert round(orders["Sales"].sum(), 2) == 2_296_919.49
assert round(orders["Profit"].sum(), 2) == 286_409.08
assert orders["Ship Date"].notna().all()
assert (pd.to_datetime(orders["Ship Date"]) >= pd.to_datetime(orders["Order Date"])).all()
assert orders["Sales"].gt(0).all()
assert orders["Quantity"].gt(0).all()
assert orders["Discount"].between(0, 1).all()
assert set(orders["Region"]) == set(managers["Region"])

wb = load_workbook(workbook_path, data_only=False)
assert expected_sheets.issubset(wb.sheetnames)
data = wb["Data"]
assert data.max_row == len(orders) + 1
assert data["AC1"].value == "First Line of Order"
assert data["AC2"].value.startswith("=IF(COUNTIF($B$2:B2")
assert data["AC3"].value.startswith("=IF(COUNTIF($B$2:B3")
assert data["AD1"].value == "Discount Band"
assert wb["Monthly Trend"]["D54"].value.startswith("=ROUND(D53-")
assert wb["Dashboard"]["B1"].value == "Superstore Sales Dashboard, 2015-2018"
print("Validation passed: cleaned data and workbook structure are consistent.")
